# -*- coding: utf-8 -*-
"""
Generate complementary prevalence analyses and Supplementary Tables S2–S5.

Analyses
--------
1. Pooled prevalence using colour-specific garment variants.
2. Portal-specific prevalence for the four retailer–market portals.
3. Distinct-product prevalence using an any-variant criterion.
4. Conservative, default, and expanded rule-boundary comparisons.

Generated Supplementary Information tables
-------------------------------------------
- Table S2: dataset-construction flow, sample structure, and analytical denominators
- Table S3: removable-hardware, retained-barrier, and overall core-disruptor prevalence across analytical boundaries
- Table S4: rule-level pooled-variant and distinct-product prevalence
- Table S5: portal-specific prevalence by indicator and rule boundary

The script also exports calculation-ready long/wide datasets, product-family
audits, agreement diagnostics, an analysis-design file, and QA results.
"""

from __future__ import annotations

import argparse
import json
import re
from collections import OrderedDict
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "python-docx is required to create the supplementary Word tables. "
        "Install it with: pip install python-docx"
    ) from exc


SCRIPT_DIR = Path(__file__).resolve().parent
LOCAL_PROJECT_DIR = Path(
    r"C:\Users\laptop-kl\OneDrive - Universiteit Leiden\PlasticTradeFlow"
    r"\data_mining_clothing\textile_preprocessing\waste_management\1st_revision"
)
PROJECT_DIR = LOCAL_PROJECT_DIR if LOCAL_PROJECT_DIR.exists() else SCRIPT_DIR

DEFAULT_JSONL = PROJECT_DIR / "6_JSONL_component_normalized.jsonl"
if not DEFAULT_JSONL.exists():
    DEFAULT_JSONL = SCRIPT_DIR / "6_JSONL_component_normalized.jsonl"
DEFAULT_FLAGS = PROJECT_DIR / "outputs_v2" / "disruptor_rule_flags_by_variant.csv"
DEFAULT_RULE_SUMMARY = PROJECT_DIR / "outputs_v2" / "disruptor_rule_summary.csv"
DEFAULT_OUTPUT_DIR = PROJECT_DIR / "tables_S2_to_S5_v2"

SETTINGS = ["conservative", "default", "expanded"]
SETTING_LABELS = {
    "conservative": "Conservative",
    "default": "Default",
    "expanded": "Expanded",
}

# Panel A of Table S2 summarizes the upstream dataset-construction stages
# documented in the archived dataset release. The final count is independently
# checked against the row-level analysis input in run_qa().
DATASET_CONSTRUCTION_FLOW = [
    {
        "Dataset stage": "Raw scraped product-page records",
        "Records (n)": 47834,
        "Role in construction": "Starting product-page records",
    },
    {
        "Dataset stage": "After minimum-information filtering",
        "Records (n)": 47570,
        "Role in construction": "Removed records lacking required material or colour information",
    },
    {
        "Dataset stage": "After colour-variant expansion and harmonization",
        "Records (n)": 51994,
        "Role in construction": "Expanded Uniqlo records and harmonized retailer schemas",
    },
    {
        "Dataset stage": "After category normalization and scope filtering",
        "Records (n)": 48244,
        "Role in construction": "Removed accessories, footwear, and unresolved categories",
    },
    {
        "Dataset stage": "Final analytical dataset",
        "Records (n)": 47522,
        "Role in construction": "After component normalization and consistency filtering",
    },
]

# Order is deliberate: aggregate outcomes first, then the individual rules and
# diagnostics. R7b and R8b remain diagnostics and are not included in aggregate
# the retained-barrier aggregate or overall core-disruptor indicator.
METRICS = OrderedDict(
    [
        ("any_removable_disruptor", "Any removable-hardware indicator (R1–R3)"),
        ("any_retained_barrier", "Any retained-barrier indicator (R4–R8)"),
        ("any_disruptor_overall", "Any core disruptor indicator (R1–R8)"),
        ("r1_metal_hardware", "R1 Metal-specified/conventional hardware"),
        ("r2_plastic_hardware", "R2 Plastic-specified hardware"),
        ("r3_mixed_hardware", "R3 Material-ambiguous hardware"),
        ("r4_fabric_attached_trim", "R4 Fabric attached trim"),
        ("r5_decorative_nontextile", "R5 Decorative/non-textile attachment"),
        ("r6_surface_print_coating", "R6 Surface print/coating barrier"),
        ("r7_lining_multilayer", "R7 Lining/multilayer presence"),
        ("r7b_hidden_layer_mismatch", "R7b Hidden-layer mismatch diagnostic"),
        ("r8_secondary_component_presence", "R8 Secondary-component presence"),
        (
            "r8b_secondary_material_difference",
            "R8b Secondary-component mismatch diagnostic",
        ),
    ]
)

AGGREGATE_METRICS = [
    "any_removable_disruptor",
    "any_retained_barrier",
    "any_disruptor_overall",
]
RULE_AND_DIAGNOSTIC_METRICS = [
    "r1_metal_hardware",
    "r2_plastic_hardware",
    "r3_mixed_hardware",
    "r4_fabric_attached_trim",
    "r5_decorative_nontextile",
    "r6_surface_print_coating",
    "r7_lining_multilayer",
    "r7b_hidden_layer_mismatch",
    "r8_secondary_component_presence",
    "r8b_secondary_material_difference",
]

PORTAL_ORDER = [
    ("hm", "au"),
    ("hm", "gb"),
    ("uniqlo", "au"),
    ("uniqlo", "uk"),
]
PORTAL_LABELS = {
    ("hm", "au"): "H&M Australia",
    ("hm", "gb"): "H&M Great Britain",
    ("uniqlo", "au"): "Uniqlo Australia",
    ("uniqlo", "uk"): "Uniqlo UK",
}

PRIMARY_SCOPE_ID = "pooled_variant"
PRODUCT_SCOPE_ID = "distinct_product_any_variant"

ANALYSIS_DESIGN = [
    {
        "analysis_name": "Pooled variant-level prevalence",
        "analysis_type": "Primary prevalence analysis",
        "unit_counted": "Colour-specific garment variant",
        "source_boundary": "All four retailer–market portals pooled",
        "calculation": "Number of flagged variants divided by all variants",
        "question_answered": (
            "What share of all observed colour-specific garment variants contains "
            "the indicator?"
        ),
        "interpretation_limit": (
            "Describes the fixed observed garment dataset; it is not a probability estimate "
            "for the global garment market."
        ),
    },
    {
        "analysis_name": "Portal-specific variant-level prevalence",
        "analysis_type": "Complementary source-boundary analysis",
        "unit_counted": "Colour-specific garment variant",
        "source_boundary": (
            "Calculated separately for H&M Australia, H&M Great Britain, "
            "Uniqlo Australia, and Uniqlo UK"
        ),
        "calculation": (
            "Within each portal, number of flagged variants divided by all variants "
            "in that portal"
        ),
        "question_answered": (
            "Is the main pattern visible within each retailer–market portal, "
            "rather than only after pooling the four sources?"
        ),
        "interpretation_limit": (
            "Portal differences are descriptive and must not be interpreted as causal "
            "brand or country effects."
        ),
    },
    {
        "analysis_name": "Distinct-product prevalence: at least one flagged variant",
        "analysis_type": "Complementary analytical-unit analysis",
        "unit_counted": "Retailer-specific product family across regions and colours",
        "source_boundary": "All portals combined after product aggregation",
        "calculation": (
            "For each indicator, a product equals 1 when at least one observed colour "
            "or regional variant is flagged; product prevalence is the mean of this "
            "product-level binary variable"
        ),
        "question_answered": (
            "What share of distinct product families has at least one observed variant "
            "containing the indicator?"
        ),
        "interpretation_limit": (
            "This is a precautionary any-variant criterion. It answers a different "
            "question from variant prevalence and does not prove product independence."
        ),
    },
    {
        "analysis_name": "Rule-boundary sensitivity",
        "analysis_type": "Sensitivity analysis",
        "unit_counted": "Applied to every prevalence perspective above",
        "source_boundary": "Conservative, default, and expanded rule definitions",
        "calculation": (
            "Recalculate every prevalence measure using each existing rule-boundary setting"
        ),
        "question_answered": (
            "How strongly do the results depend on narrower or broader operational rule definitions?"
        ),
        "interpretation_limit": (
            "Tests observable retailer-data rule boundaries; it does not physically validate "
            "garment construction or process compatibility."
        ),
    },
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--jsonl", type=Path, default=DEFAULT_JSONL)
    parser.add_argument("--flags-csv", type=Path, default=DEFAULT_FLAGS)
    parser.add_argument("--rule-summary-csv", type=Path, default=DEFAULT_RULE_SUMMARY)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    return parser.parse_args()


def metric_columns() -> list[str]:
    return [f"{metric}_{setting}" for metric in METRICS for setting in SETTINGS]


def validate_input(df: pd.DataFrame) -> None:
    required_id = {
        "row_id",
        "brand",
        "region",
        "parent_product_id",
        "product_name",
        "detail_category",
        "variant_colour",
    }
    missing = sorted((required_id | set(metric_columns())) - set(df.columns))
    if missing:
        raise ValueError(f"Missing required input columns: {missing}")

    if df.empty:
        raise ValueError("The row-level flag file is empty.")

    if df["parent_product_id"].isna().any():
        raise ValueError("parent_product_id contains missing values.")

    hm_ids = df.loc[df["brand"] == "hm", "parent_product_id"]
    uniqlo_ids = df.loc[df["brand"] == "uniqlo", "parent_product_id"]
    if not hm_ids.str.fullmatch(r"\d{10}").all():
        raise ValueError("All H&M parent_product_id values must be ten-digit article identifiers.")
    if not uniqlo_ids.str.fullmatch(r"\d+").all():
        raise ValueError("All Uniqlo parent_product_id values must be numeric identifiers.")

    observed_portals = set(map(tuple, df[["brand", "region"]].drop_duplicates().values))
    unexpected = observed_portals - set(PORTAL_ORDER)
    missing_portals = set(PORTAL_ORDER) - observed_portals
    if unexpected or missing_portals:
        raise ValueError(
            f"Portal mismatch. Unexpected={sorted(unexpected)}; missing={sorted(missing_portals)}"
        )

    bad_columns: list[str] = []
    for col in metric_columns():
        values = set(pd.to_numeric(df[col], errors="coerce").dropna().unique())
        if not values.issubset({0, 1}):
            bad_columns.append(col)
    if bad_columns:
        raise ValueError(f"Flag columns contain values other than 0/1: {bad_columns}")


def attach_source_urls(df: pd.DataFrame, jsonl_path: Path) -> pd.DataFrame:
    """Attach source URLs by the evaluator's one-based JSONL line number."""
    rows: list[dict[str, object]] = []
    with jsonl_path.open("r", encoding="utf-8") as handle:
        for row_id, line in enumerate(handle, start=1):
            record = json.loads(line)
            rows.append(
                {
                    "row_id": row_id,
                    "source_url": record.get("url"),
                    "source_brand": record.get("brand"),
                    "source_region": record.get("region"),
                    "source_parent_product_id": str(record.get("parent_product_id")),
                }
            )
    source = pd.DataFrame(rows)
    merged = df.merge(source, on="row_id", how="left", validate="one_to_one")
    if merged["source_url"].isna().any():
        raise ValueError("Some row-level flags could not be linked to JSONL source URLs.")
    identity_ok = (
        merged["brand"].astype(str).eq(merged["source_brand"].astype(str))
        & merged["region"].astype(str).eq(merged["source_region"].astype(str))
        & merged["parent_product_id"].astype(str).eq(
            merged["source_parent_product_id"].astype(str)
        )
    )
    if not identity_ok.all():
        raise ValueError("Row order or identifiers differ between the flag CSV and JSONL input.")
    return merged.drop(
        columns=["source_brand", "source_region", "source_parent_product_id"]
    )


def derive_product_family_id(df: pd.DataFrame) -> pd.DataFrame:
    """Create a retailer-consistent product-family identifier.

    H&M uses ten-digit article identifiers in which the first seven digits
    identify the style/product family and the final three digits identify the
    colour article. Uniqlo colour variants expanded from one product page share
    the full URL product-page identifier (for example, E484929-000); the suffix
    must be retained because different suffixes can represent distinct product
    pages even when the six-digit base identifier is the same.
    """
    result = df.copy()
    result["analysis_product_id"] = pd.Series(pd.NA, index=result.index, dtype="string")

    hm_mask = result["brand"].eq("hm")
    result.loc[hm_mask, "analysis_product_id"] = result.loc[
        hm_mask, "parent_product_id"
    ].str.slice(0, 7)

    uniqlo_mask = result["brand"].eq("uniqlo")
    uniqlo_key = result.loc[uniqlo_mask, "source_url"].str.extract(
        r"/products/(E\d+-\d+)", expand=False
    )
    result.loc[uniqlo_mask, "analysis_product_id"] = uniqlo_key

    if result["analysis_product_id"].isna().any():
        bad = result.loc[
            result["analysis_product_id"].isna(), ["brand", "source_url"]
        ].head(10)
        raise ValueError(f"Could not derive product-family identifiers:\n{bad}")
    return result

def audit_product_family_key(df: pd.DataFrame) -> pd.DataFrame:
    """Document internal consistency of the derived product-family key."""
    rows: list[dict[str, object]] = []
    for brand, region in PORTAL_ORDER:
        subset = df[(df["brand"] == brand) & (df["region"] == region)]
        grouped = subset.groupby(["brand", "region", "analysis_product_id"], dropna=False)
        name_counts = grouped["product_name"].nunique()
        category_counts = grouped["detail_category"].nunique()
        rows.append(
            {
                "portal": PORTAL_LABELS[(brand, region)],
                "brand": brand,
                "region": region,
                "product_key_rule": (
                    "First seven digits of ten-digit H&M article ID"
                    if brand == "hm"
                    else "Full Uniqlo URL product-page identifier (E######-###)"
                ),
                "n_product_families": len(name_counts),
                "n_families_with_multiple_product_names_within_portal": int((name_counts > 1).sum()),
                "n_families_with_multiple_detail_categories_within_portal": int((category_counts > 1).sum()),
                "share_families_single_name_percent": float((name_counts == 1).mean() * 100),
                "share_families_single_category_percent": float((category_counts == 1).mean() * 100),
            }
        )
    return pd.DataFrame(rows)


def prevalence_records(
    frame: pd.DataFrame,
    analysis_id: str,
    analysis_label: str,
    analysis_family: str,
    unit_of_analysis: str,
    source_boundary: str,
    aggregation_rule: str,
) -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    n_units = len(frame)
    for metric, metric_label in METRICS.items():
        for setting in SETTINGS:
            col = f"{metric}_{setting}"
            n_flagged = int(frame[col].sum())
            prevalence = n_flagged / n_units if n_units else np.nan
            records.append(
                {
                    "analysis_id": analysis_id,
                    "analysis_label": analysis_label,
                    "analysis_family": analysis_family,
                    "unit_of_analysis": unit_of_analysis,
                    "source_boundary": source_boundary,
                    "aggregation_rule": aggregation_rule,
                    "metric": metric,
                    "metric_label": metric_label,
                    "setting": setting,
                    "setting_label": SETTING_LABELS[setting],
                    "n_units": n_units,
                    "n_flagged": n_flagged,
                    "prevalence": prevalence,
                    "prevalence_percent": prevalence * 100,
                }
            )
    return records


def calculate_prevalence(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    all_records: list[dict[str, object]] = []

    all_records.extend(
        prevalence_records(
            df,
            analysis_id=PRIMARY_SCOPE_ID,
            analysis_label="Pooled variant-level prevalence",
            analysis_family="Primary prevalence analysis",
            unit_of_analysis="Colour-specific garment variant",
            source_boundary="All four retailer–market portals pooled",
            aggregation_rule="Direct binary flag for each variant",
        )
    )

    for brand, region in PORTAL_ORDER:
        portal = df[(df["brand"] == brand) & (df["region"] == region)].copy()
        portal_label = PORTAL_LABELS[(brand, region)]
        all_records.extend(
            prevalence_records(
                portal,
                analysis_id=f"portal_variant_{brand}_{region}",
                analysis_label=f"{portal_label}: variant-level prevalence",
                analysis_family="Complementary source-boundary analysis",
                unit_of_analysis="Colour-specific garment variant",
                source_boundary=portal_label,
                aggregation_rule="Direct binary flag for each variant within portal",
            )
        )

    # A distinct product family collapses all observed colours and both regional
    # portals using the retailer-consistent analysis_product_id. The any-variant
    # criterion uses max() because the input flags are binary.
    products = (
        df.groupby(["brand", "analysis_product_id"], as_index=False)[metric_columns()]
        .max()
        .sort_values(["brand", "analysis_product_id"])
        .reset_index(drop=True)
    )
    all_records.extend(
        prevalence_records(
            products,
            analysis_id=PRODUCT_SCOPE_ID,
            analysis_label="Distinct-product prevalence: at least one flagged variant",
            analysis_family="Complementary analytical-unit analysis",
            unit_of_analysis="Retailer-specific product family",
            source_boundary="All portals combined after product aggregation",
            aggregation_rule=(
                "Product flagged when at least one observed colour or regional variant is flagged"
            ),
        )
    )

    long_df = pd.DataFrame(all_records)
    return long_df, products


def calculate_sample_structure(df: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []

    def summarize(
        subset: pd.DataFrame,
        boundary_label: str,
        product_keys: list[str],
        boundary_type: str,
    ) -> dict[str, object]:
        multiplicity = subset.groupby(product_keys, dropna=False).size()
        return {
            "boundary_type": boundary_type,
            "boundary": boundary_label,
            "n_colour_variants": len(subset),
            "n_product_units": int(len(multiplicity)),
            "n_products_with_multiple_variants": int((multiplicity > 1).sum()),
            "share_products_with_multiple_variants": float((multiplicity > 1).mean()),
            "mean_variants_per_product": float(multiplicity.mean()),
            "median_variants_per_product": float(multiplicity.median()),
            "p90_variants_per_product": float(multiplicity.quantile(0.90)),
            "max_variants_per_product": int(multiplicity.max()),
        }

    rows.append(
        summarize(
            df,
            "All portals pooled: distinct products",
            ["brand", "analysis_product_id"],
            "Pooled cross-market product-family boundary",
        )
    )
    rows.append(
        summarize(
            df,
            "All portals pooled: product-market units",
            ["brand", "region", "analysis_product_id"],
            "Pooled portal-specific product-family boundary",
        )
    )
    for brand, region in PORTAL_ORDER:
        subset = df[(df["brand"] == brand) & (df["region"] == region)]
        rows.append(
            summarize(
                subset,
                PORTAL_LABELS[(brand, region)],
                ["brand", "region", "analysis_product_id"],
                "Retailer-market portal",
            )
        )

    result = pd.DataFrame(rows)
    result["share_products_with_multiple_variants_percent"] = (
        result["share_products_with_multiple_variants"] * 100
    )
    return result


def calculate_within_product_agreement(df: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    grouped = df.groupby(["brand", "analysis_product_id"], dropna=False)

    for metric, metric_label in METRICS.items():
        for setting in SETTINGS:
            col = f"{metric}_{setting}"
            stats = grouped[col].agg(["min", "max", "size"])
            mixed = stats["min"] != stats["max"]
            multi = stats["size"] > 1
            all_flagged = stats["min"] == 1
            all_unflagged = stats["max"] == 0
            rows.append(
                {
                    "metric": metric,
                    "metric_label": metric_label,
                    "setting": setting,
                    "setting_label": SETTING_LABELS[setting],
                    "n_distinct_products": len(stats),
                    "n_all_variants_unflagged": int(all_unflagged.sum()),
                    "n_all_variants_flagged": int(all_flagged.sum()),
                    "n_mixed_variant_status": int(mixed.sum()),
                    "share_mixed_all_products": float(mixed.mean()),
                    "share_mixed_all_products_percent": float(mixed.mean() * 100),
                    "n_multi_variant_products": int(multi.sum()),
                    "n_mixed_among_multi_variant_products": int((mixed & multi).sum()),
                    "share_mixed_among_multi_variant_products": (
                        float(mixed[multi].mean()) if multi.any() else np.nan
                    ),
                    "share_mixed_among_multi_variant_products_percent": (
                        float(mixed[multi].mean() * 100) if multi.any() else np.nan
                    ),
                }
            )
    return pd.DataFrame(rows)


def calculate_cross_market_overlap(df: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    region_counts = (
        df.groupby(["brand", "analysis_product_id"], dropna=False)["region"]
        .nunique()
        .rename("n_regions")
        .reset_index()
    )
    overlap_rows: list[dict[str, object]] = []
    for brand_label, subset in [("All brands", region_counts)] + [
        (brand, region_counts[region_counts["brand"] == brand])
        for brand in sorted(region_counts["brand"].unique())
    ]:
        overlap_rows.append(
            {
                "brand": brand_label,
                "n_distinct_products": len(subset),
                "n_products_in_both_regions": int((subset["n_regions"] > 1).sum()),
                "share_products_in_both_regions": float((subset["n_regions"] > 1).mean()),
                "share_products_in_both_regions_percent": float(
                    (subset["n_regions"] > 1).mean() * 100
                ),
            }
        )
    overlap_summary = pd.DataFrame(overlap_rows)

    # First collapse colours within each portal using the same any-variant rule.
    portal_products = (
        df.groupby(["brand", "region", "analysis_product_id"], as_index=False)[
            metric_columns()
        ]
        .max()
    )
    repeated_ids = (
        portal_products.groupby(["brand", "analysis_product_id"])["region"]
        .nunique()
        .loc[lambda x: x > 1]
        .index
    )
    repeated = portal_products.set_index(["brand", "analysis_product_id"]).loc[
        repeated_ids
    ].reset_index()

    agreement_rows: list[dict[str, object]] = []
    repeated_grouped = repeated.groupby(["brand", "analysis_product_id"])
    for metric, metric_label in METRICS.items():
        for setting in SETTINGS:
            col = f"{metric}_{setting}"
            stats = repeated_grouped[col].agg(["min", "max"])
            agree = stats["min"] == stats["max"]
            agreement_rows.append(
                {
                    "metric": metric,
                    "metric_label": metric_label,
                    "setting": setting,
                    "setting_label": SETTING_LABELS[setting],
                    "n_cross_market_products": len(stats),
                    "n_same_status_across_regions": int(agree.sum()),
                    "n_different_status_across_regions": int((~agree).sum()),
                    "agreement_share": float(agree.mean()),
                    "agreement_percent": float(agree.mean() * 100),
                }
            )
    agreement = pd.DataFrame(agreement_rows)
    return overlap_summary, agreement


def make_wide_prevalence(long_df: pd.DataFrame) -> pd.DataFrame:
    temp = long_df.copy()
    temp["scope_setting"] = (
        temp["analysis_id"] + "__" + temp["setting"]
    )
    wide = temp.pivot(
        index=["metric", "metric_label"],
        columns="scope_setting",
        values="prevalence_percent",
    ).reset_index()
    wide.columns.name = None
    metric_order = {metric: i for i, metric in enumerate(METRICS)}
    wide["_order"] = wide["metric"].map(metric_order)
    return wide.sort_values("_order").drop(columns="_order").reset_index(drop=True)


def lookup_value(
    long_df: pd.DataFrame,
    analysis_id: str,
    metric: str,
    setting: str,
    field: str = "prevalence_percent",
) -> float:
    matched = long_df[
        (long_df["analysis_id"] == analysis_id)
        & (long_df["metric"] == metric)
        & (long_df["setting"] == setting)
    ]
    if len(matched) != 1:
        raise ValueError(
            f"Expected one row for {analysis_id}, {metric}, {setting}; found {len(matched)}"
        )
    return float(matched.iloc[0][field])


def build_supplementary_tables(
    long_df: pd.DataFrame,
    sample_structure: pd.DataFrame,
) -> dict[str, pd.DataFrame]:
    # Table S2, Panel A: upstream dataset-construction flow documented in the
    # archived dataset release. Panel B: analytical denominators and colour-variant
    # multiplicity calculated directly from the final analysis input.
    s2a = pd.DataFrame(DATASET_CONSTRUCTION_FLOW)

    # Panel B uses the distinct-product pooled boundary and the four portals.
    # The product–market pooled row remains in the audit file but is omitted from
    # the publication-facing table to avoid unnecessary complexity.
    selected_boundaries = [
        "All portals pooled: distinct products",
        *[PORTAL_LABELS[p] for p in PORTAL_ORDER],
    ]
    s2b = sample_structure[
        sample_structure["boundary"].isin(selected_boundaries)
    ].copy()
    s2b["_order"] = s2b["boundary"].map(
        {boundary: i for i, boundary in enumerate(selected_boundaries)}
    )
    s2b = s2b.sort_values("_order").drop(columns="_order")
    s2b = s2b[
        [
            "boundary",
            "n_colour_variants",
            "n_product_units",
            "n_products_with_multiple_variants",
            "share_products_with_multiple_variants_percent",
            "mean_variants_per_product",
            "median_variants_per_product",
            "max_variants_per_product",
        ]
    ].rename(
        columns={
            "boundary": "Analytical boundary",
            "n_colour_variants": "Colour variants (n)",
            "n_product_units": "Product units (n)",
            "n_products_with_multiple_variants": "Products with >1 variant (n)",
            "share_products_with_multiple_variants_percent": "Products with >1 variant (%)",
            "mean_variants_per_product": "Mean variants/product",
            "median_variants_per_product": "Median variants/product",
            "max_variants_per_product": "Maximum variants/product",
        }
    )

    # Table S3: aggregate outcomes for every prevalence perspective and setting.
    scope_order = [
        PRIMARY_SCOPE_ID,
        *[f"portal_variant_{brand}_{region}" for brand, region in PORTAL_ORDER],
        PRODUCT_SCOPE_ID,
    ]
    scope_labels = {
        PRIMARY_SCOPE_ID: "All portals pooled: colour variants",
        PRODUCT_SCOPE_ID: "Distinct products: at least one flagged variant",
        **{
            f"portal_variant_{brand}_{region}": PORTAL_LABELS[(brand, region)]
            for brand, region in PORTAL_ORDER
        },
    }
    s3_rows: list[dict[str, object]] = []
    for scope in scope_order:
        row: dict[str, object] = {"Analytical boundary": scope_labels[scope]}
        n_units = int(
            lookup_value(long_df, scope, "any_disruptor_overall", "default", "n_units")
        )
        row["Units (n)"] = n_units
        for metric in AGGREGATE_METRICS:
            short = {
                "any_removable_disruptor": "Any removable",
                "any_retained_barrier": "Any retained",
                "any_disruptor_overall": "Any overall",
            }[metric]
            for setting in SETTINGS:
                row[f"{short} - {SETTING_LABELS[setting]}"] = lookup_value(
                    long_df, scope, metric, setting
                )
        s3_rows.append(row)
    s3 = pd.DataFrame(s3_rows)

    # Table S4: rule and diagnostic prevalence for the primary and product unit.
    s4_rows: list[dict[str, object]] = []
    for metric in RULE_AND_DIAGNOSTIC_METRICS:
        row: dict[str, object] = {"Indicator": METRICS[metric]}
        for scope, scope_short in [
            (PRIMARY_SCOPE_ID, "Pooled variants"),
            (PRODUCT_SCOPE_ID, "Distinct products"),
        ]:
            for setting in SETTINGS:
                row[f"{scope_short} - {SETTING_LABELS[setting]}"] = lookup_value(
                    long_df, scope, metric, setting
                )
        s4_rows.append(row)
    s4 = pd.DataFrame(s4_rows)

    # Table S5: portal-specific values for all indicators and settings. It is
    # structured long to remain readable in Word and easy to inspect.
    s5_rows: list[dict[str, object]] = []
    for metric, metric_label in METRICS.items():
        for setting in SETTINGS:
            row: dict[str, object] = {
                "Indicator": metric_label,
                "Rule boundary": SETTING_LABELS[setting],
            }
            row["All portals pooled"] = lookup_value(
                long_df, PRIMARY_SCOPE_ID, metric, setting
            )
            for brand, region in PORTAL_ORDER:
                scope = f"portal_variant_{brand}_{region}"
                row[PORTAL_LABELS[(brand, region)]] = lookup_value(
                    long_df, scope, metric, setting
                )
            s5_rows.append(row)
    s5 = pd.DataFrame(s5_rows)

    return {"S2A": s2a, "S2B": s2b, "S3": s3, "S4": s4, "S5": s5}


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, font_size: float = 8.0) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(font_size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def format_table_value(column: str, value: object) -> str:
    if pd.isna(value):
        return "-"
    if "(%)" in column or column.startswith(("Any ", "Pooled variants", "Distinct products")):
        return f"{float(value):.1f}"
    if column in {
        "Records (n)",
        "Colour variants (n)",
        "Product units (n)",
        "Products with >1 variant (n)",
        "Maximum variants/product",
        "Units (n)",
    }:
        return f"{int(round(float(value))):,}"
    if column in {"Mean variants/product"}:
        return f"{float(value):.2f}"
    if column in {"Median variants/product"}:
        return f"{float(value):.1f}"
    if isinstance(value, (float, np.floating)):
        return f"{float(value):.1f}"
    return str(value)


def add_dataframe_table(
    document: Document,
    dataframe: pd.DataFrame,
    caption: str,
    note: str,
    font_size: float = 7.5,
) -> None:
    caption_p = document.add_paragraph()
    caption_p.paragraph_format.space_before = Pt(6)
    caption_p.paragraph_format.space_after = Pt(4)
    caption_run = caption_p.add_run(caption)
    caption_run.bold = True
    caption_run.font.name = "Times New Roman"
    caption_run.font.size = Pt(9)

    table = document.add_table(rows=1, cols=len(dataframe.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, column in enumerate(dataframe.columns):
        set_cell_text(header.cells[idx], column, bold=True, font_size=font_size)

    for _, data_row in dataframe.iterrows():
        cells = table.add_row().cells
        for idx, column in enumerate(dataframe.columns):
            text = format_table_value(column, data_row[column])
            set_cell_text(cells[idx], text, bold=False, font_size=font_size)
            if idx == 0:
                cells[idx].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT

    note_p = document.add_paragraph()
    note_p.paragraph_format.space_before = Pt(3)
    note_p.paragraph_format.space_after = Pt(8)
    note_run = note_p.add_run(f"Note: {note}")
    note_run.italic = True
    note_run.font.name = "Times New Roman"
    note_run.font.size = Pt(8)


def create_supplementary_docx(
    tables: dict[str, pd.DataFrame],
    output_path: Path,
) -> None:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Supplementary Tables S2–S5")
    title_run.bold = True
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(13)

    intro = doc.add_paragraph(
        "These tables distinguish complementary prevalence analyses based on "
        "alternative analytical units and retailer–market boundaries from the "
        "conservative/default/expanded rule-boundary sensitivity analysis."
    )
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    table_s2_caption = doc.add_paragraph()
    table_s2_caption.paragraph_format.space_before = Pt(6)
    table_s2_caption.paragraph_format.space_after = Pt(4)
    table_s2_run = table_s2_caption.add_run(
        "Table S2. Dataset-construction flow, sample structure, and analytical "
        "denominators used in the prevalence analyses."
    )
    table_s2_run.bold = True
    table_s2_run.font.name = "Times New Roman"
    table_s2_run.font.size = Pt(9)

    add_dataframe_table(
        doc,
        tables["S2A"],
        "Panel A. Dataset-construction flow",
        (
            "The first four counts summarize the upstream construction workflow documented "
            "in the archived dataset release. The final analytical count is checked against "
            "the row-level rule-evaluation input."
        ),
        font_size=7.5,
    )

    add_dataframe_table(
        doc,
        tables["S2B"],
        "Panel B. Analytical denominators and colour-variant multiplicity",
        (
            "For H&M, the product-family key is the seven-digit style prefix of the ten-digit "
            "article identifier; for Uniqlo, it is the full product-page identifier including "
            "its suffix. The same retailer-specific key is used across both regions. Garment "
            "categories were assigned after product acquisition and were not used as sampling quotas."
        ),
        font_size=7.2,
    )

    add_dataframe_table(
        doc,
        tables["S3"],
        (
            "Table S3. Aggregate removable-hardware, retained-barrier, and overall core-disruptor prevalence across complementary analytical units "
            "and retailer–market boundaries (%)."
        ),
        (
            "Pooled and portal-specific estimates use colour-specific garment variants as the unit. "
            "Distinct-product estimates count each retailer-specific product family once and "
            "classify it as flagged when at least one observed colour or regional variant is flagged. "
            "Conservative, default, and expanded refer to rule-boundary settings. "
            "‘Any removable’ denotes at least one removable-hardware indicator in R1–R3; "
            "‘Any retained’ denotes at least one retained-barrier indicator in R4–R8; and "
            "‘Any overall’ denotes at least one core disruptor indicator in R1–R8."
        ),
        font_size=6.8,
    )

    add_dataframe_table(
        doc,
        tables["S4"],
        (
            "Table S4. Rule-level prevalence under pooled variant-level and distinct-product "
            "analytical units (%)."
        ),
        (
            "R7b and R8b are material-mismatch diagnostics and are not included in the "
            "retained-barrier aggregate or the overall core-disruptor indicator. The distinct-product measure uses "
            "an at-least-one-flagged-variant criterion."
        ),
        font_size=7.2,
    )

    add_dataframe_table(
        doc,
        tables["S5"],
        (
            "Table S5. Pooled and portal-specific variant-level prevalence by indicator and rule boundary (%)."
        ),
        (
            "All columns use colour-specific garment variants as the unit. The pooled column "
            "reproduces the primary main-manuscript estimate and is included as a reference for "
            "the four portal-specific estimates. Portal differences are descriptive and should "
            "not be interpreted as causal retailer or country effects."
        ),
        font_size=6.8,
    )

    doc.core_properties.title = "Supplementary Tables S2–S5"
    doc.core_properties.subject = "Complementary prevalence analyses"
    doc.save(output_path)


def run_qa(
    df: pd.DataFrame,
    long_df: pd.DataFrame,
    products: pd.DataFrame,
    rule_summary_path: Path,
    supplementary_tables: dict[str, pd.DataFrame],
    product_key_audit: pd.DataFrame,
) -> pd.DataFrame:
    checks: list[dict[str, object]] = []

    def add(name: str, passed: bool, observed: object, expected: object, detail: str) -> None:
        checks.append(
            {
                "check": name,
                "passed": bool(passed),
                "observed": observed,
                "expected": expected,
                "detail": detail,
            }
        )

    add(
        "variant_row_count",
        len(df) == 47522,
        len(df),
        47522,
        "Primary analytical dataset size.",
    )
    add(
        "distinct_product_count",
        len(products) == df.groupby(["brand", "analysis_product_id"]).ngroups,
        len(products),
        df.groupby(["brand", "analysis_product_id"]).ngroups,
        "One row per retailer-specific product family.",
    )
    key_consistency = bool(
        (product_key_audit["n_families_with_multiple_product_names_within_portal"] == 0).all()
        and (product_key_audit["n_families_with_multiple_detail_categories_within_portal"] == 0).all()
    )
    add(
        "product_family_key_internal_consistency",
        key_consistency,
        key_consistency,
        True,
        "Within each portal, every derived product-family key maps to one product name and one detail category.",
    )
    portal_total = sum(
        int(
            long_df[
                (long_df["analysis_id"] == f"portal_variant_{b}_{r}")
                & (long_df["metric"] == "any_disruptor_overall")
                & (long_df["setting"] == "default")
            ]["n_units"].iloc[0]
        )
        for b, r in PORTAL_ORDER
    )
    add(
        "portal_denominators_sum_to_pooled",
        portal_total == len(df),
        portal_total,
        len(df),
        "The four portal strata partition the pooled variant dataset.",
    )

    if rule_summary_path.exists():
        baseline = pd.read_csv(rule_summary_path)
        baseline_lookup = baseline.set_index("metric")["share_flagged"]
        max_difference = 0.0
        missing_metrics: list[str] = []
        for metric in METRICS:
            for setting in SETTINGS:
                source_metric = f"{metric}_{setting}"
                if source_metric not in baseline_lookup.index:
                    missing_metrics.append(source_metric)
                    continue
                observed = lookup_value(
                    long_df,
                    PRIMARY_SCOPE_ID,
                    metric,
                    setting,
                    "prevalence",
                )
                max_difference = max(
                    max_difference,
                    abs(observed - float(baseline_lookup[source_metric])),
                )
        add(
            "pooled_results_match_rule_summary",
            (not missing_metrics) and max_difference < 1e-12,
            max_difference,
            "< 1e-12",
            f"Missing metrics: {missing_metrics or 'none'}",
        )
    else:
        add(
            "pooled_results_match_rule_summary",
            False,
            "not checked",
            str(rule_summary_path),
            "Rule-summary CSV was not found.",
        )

    # Product-level max aggregation should equal a direct any() calculation.
    direct = (
        df.groupby(["brand", "analysis_product_id"])[metric_columns()]
        .any()
        .astype(int)
        .reset_index()
        .sort_values(["brand", "analysis_product_id"])
        .reset_index(drop=True)
    )
    product_sorted = (
        products.sort_values(["brand", "analysis_product_id"])
        .reset_index(drop=True)
    )

    # DataFrame.equals() is intentionally not used here because it also
    # requires identical pandas dtypes. Depending on the pandas version and
    # CSV type inference, groupby.max() may retain bool/int8/int64 columns
    # while groupby.any().astype(int) uses a different integer dtype, even
    # though the underlying binary values are identical. Normalize keys and
    # flags before comparing the actual analytical result.
    direct_keys = direct[["brand", "analysis_product_id"]].astype("string")
    product_keys = product_sorted[["brand", "analysis_product_id"]].astype("string")
    keys_equal = bool(
        np.array_equal(
            direct_keys.to_numpy(dtype=str),
            product_keys.to_numpy(dtype=str),
        )
    )

    direct_flags = (
        direct[metric_columns()]
        .apply(pd.to_numeric, errors="raise")
        .astype("int8")
    )
    product_flags = (
        product_sorted[metric_columns()]
        .apply(pd.to_numeric, errors="raise")
        .astype("int8")
    )
    mismatch_cells = int(
        np.count_nonzero(
            direct_flags.to_numpy() != product_flags.to_numpy()
        )
    )
    product_equal = keys_equal and mismatch_cells == 0

    add(
        "product_any_variant_aggregation",
        product_equal,
        f"keys_equal={keys_equal}; mismatched_flag_cells={mismatch_cells}",
        "keys_equal=True; mismatched_flag_cells=0",
        (
            "Normalized-value comparison confirms that groupby.max() and "
            "groupby.any() produce the same product-level binary flags; "
            "harmless pandas dtype differences are ignored."
        ),
    )

    expected_long_rows = (1 + len(PORTAL_ORDER) + 1) * len(METRICS) * len(SETTINGS)
    add(
        "complete_prevalence_matrix",
        len(long_df) == expected_long_rows,
        len(long_df),
        expected_long_rows,
        "Six analytical boundaries x thirteen outcomes x three settings.",
    )

    no_missing = not long_df[
        ["n_units", "n_flagged", "prevalence", "prevalence_percent"]
    ].isna().any().any()
    add(
        "no_missing_prevalence_values",
        no_missing,
        no_missing,
        True,
        "All calculation cells are populated.",
    )

    table_rows = {
        key: len(value) for key, value in supplementary_tables.items()
    }
    expected_table_rows = {"S2A": 5, "S2B": 5, "S3": 6, "S4": 10, "S5": 39}
    add(
        "supplementary_table_dimensions",
        table_rows == expected_table_rows,
        table_rows,
        expected_table_rows,
        "Expected publication-facing table row counts.",
    )
    final_flow_count = int(
        supplementary_tables["S2A"].loc[
            supplementary_tables["S2A"]["Dataset stage"].eq("Final analytical dataset"),
            "Records (n)",
        ].iloc[0]
    )
    add(
        "table_S2_panel_A_final_count_matches_input",
        final_flow_count == len(df),
        final_flow_count,
        len(df),
        "The final dataset-construction count matches the row-level analysis input.",
    )

    return pd.DataFrame(checks)


def write_summary(
    long_df: pd.DataFrame,
    sample_structure: pd.DataFrame,
    agreement: pd.DataFrame,
    cross_market_overlap: pd.DataFrame,
    output_path: Path,
) -> None:
    lines: list[str] = []
    lines.append("COMPLEMENTARY PREVALENCE ANALYSIS SUMMARY")
    lines.append("=" * 44)
    lines.append("")
    lines.append("Analytical structure")
    lines.append("--------------------")
    lines.append("1. Primary: pooled colour-variant prevalence.")
    lines.append("2. Complementary source boundary: portal-specific colour-variant prevalence.")
    lines.append("3. Complementary analytical unit: distinct-product prevalence using an any-variant criterion.")
    lines.append("4. Sensitivity axis: conservative/default/expanded rule boundaries.")
    lines.append("")

    pooled_n = int(
        lookup_value(
            long_df,
            PRIMARY_SCOPE_ID,
            "any_disruptor_overall",
            "default",
            "n_units",
        )
    )
    product_n = int(
        lookup_value(
            long_df,
            PRODUCT_SCOPE_ID,
            "any_disruptor_overall",
            "default",
            "n_units",
        )
    )
    lines.append("Sample structure")
    lines.append("----------------")
    lines.append(f"Colour-specific variants: {pooled_n:,}")
    lines.append(f"Distinct retailer-specific product families: {product_n:,}")
    pooled_structure = sample_structure[
        sample_structure["boundary"] == "All portals pooled: distinct products"
    ].iloc[0]
    lines.append(
        "Products with more than one observed colour/regional variant: "
        f"{int(pooled_structure['n_products_with_multiple_variants']):,} "
        f"({pooled_structure['share_products_with_multiple_variants_percent']:.1f}%)"
    )
    overall_overlap = cross_market_overlap[
        cross_market_overlap["brand"] == "All brands"
    ].iloc[0]
    lines.append(
        "Distinct product families appearing in both regional portals: "
        f"{int(overall_overlap['n_products_in_both_regions']):,} "
        f"({overall_overlap['share_products_in_both_regions_percent']:.1f}%)"
    )
    lines.append("")

    lines.append("Headline aggregate prevalence (%)")
    lines.append("---------------------------------")
    for metric in AGGREGATE_METRICS:
        lines.append(METRICS[metric])
        pooled_values = [
            lookup_value(long_df, PRIMARY_SCOPE_ID, metric, setting)
            for setting in SETTINGS
        ]
        product_values = [
            lookup_value(long_df, PRODUCT_SCOPE_ID, metric, setting)
            for setting in SETTINGS
        ]
        lines.append(
            "  Pooled variants (conservative/default/expanded): "
            + " / ".join(f"{x:.1f}" for x in pooled_values)
        )
        lines.append(
            "  Distinct products, >=1 flagged variant: "
            + " / ".join(f"{x:.1f}" for x in product_values)
        )
    lines.append("")

    portal_default = []
    for brand, region in PORTAL_ORDER:
        portal_default.append(
            (
                PORTAL_LABELS[(brand, region)],
                lookup_value(
                    long_df,
                    f"portal_variant_{brand}_{region}",
                    "any_disruptor_overall",
                    "default",
                ),
            )
        )
    lines.append("Portal-specific default overall prevalence")
    lines.append("------------------------------------------")
    for label, value in portal_default:
        lines.append(f"{label}: {value:.1f}%")
    lines.append("")

    lines.append("Within-product agreement")
    lines.append("------------------------")
    for metric in AGGREGATE_METRICS:
        row = agreement[
            (agreement["metric"] == metric) & (agreement["setting"] == "default")
        ].iloc[0]
        lines.append(
            f"{METRICS[metric]}: {int(row['n_mixed_variant_status']):,} products "
            f"have mixed variant status ({row['share_mixed_all_products_percent']:.2f}% "
            "of all distinct products; "
            f"{row['share_mixed_among_multi_variant_products_percent']:.2f}% of products "
            "with multiple observed variants)."
        )
    lines.append("")

    pooled_default = lookup_value(
        long_df, PRIMARY_SCOPE_ID, "any_disruptor_overall", "default"
    )
    product_default = lookup_value(
        long_df, PRODUCT_SCOPE_ID, "any_disruptor_overall", "default"
    )
    lines.append("Auditing interpretation")
    lines.append("----------------------")
    lines.append(
        f"The default overall estimate is {pooled_default:.2f}% at the pooled variant level "
        f"and {product_default:.2f}% at the distinct-product any-variant level."
    )
    lines.append(
        "These values answer different questions and should be reported as complementary "
        "prevalence estimates, not as competing estimates of one identical population parameter."
    )
    lines.append(
        "Portal-specific estimates show the result within each retailer–market source. "
        "They do not establish representativeness beyond the four observed portals."
    )

    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    args = parse_args()
    required_inputs = [args.jsonl, args.flags_csv, args.rule_summary_csv]
    missing_inputs = [str(path) for path in required_inputs if not path.exists()]
    if missing_inputs:
        raise FileNotFoundError(
            "Missing analysis input files:\n- " + "\n- ".join(missing_inputs)
        )

    args.output_dir.mkdir(parents=True, exist_ok=True)

    df = pd.read_csv(args.flags_csv, dtype={"parent_product_id": "string"})
    validate_input(df)
    df = attach_source_urls(df, args.jsonl)
    df = derive_product_family_id(df)
    product_key_audit = audit_product_family_key(df)

    long_df, products = calculate_prevalence(df)
    sample_structure = calculate_sample_structure(df)
    within_product = calculate_within_product_agreement(df)
    cross_overlap, cross_agreement = calculate_cross_market_overlap(df)
    wide_df = make_wide_prevalence(long_df)
    supplementary_tables = build_supplementary_tables(long_df, sample_structure)

    # Reproducibility and audit outputs.
    pd.DataFrame(ANALYSIS_DESIGN).to_csv(
        args.output_dir / "analysis_design.csv", index=False
    )
    long_df.to_csv(args.output_dir / "complementary_prevalence_long.csv", index=False)
    wide_df.to_csv(args.output_dir / "complementary_prevalence_wide.csv", index=False)
    sample_structure.to_csv(args.output_dir / "sample_structure_summary.csv", index=False)
    within_product.to_csv(args.output_dir / "within_product_agreement.csv", index=False)
    cross_overlap.to_csv(args.output_dir / "cross_market_product_family_overlap.csv", index=False)
    cross_agreement.to_csv(
        args.output_dir / "cross_market_product_family_agreement.csv", index=False
    )
    product_key_audit.to_csv(
        args.output_dir / "product_family_key_audit.csv", index=False
    )

    # Supplementary Tables S2–S5.
    supplementary_tables["S2A"].to_csv(
        args.output_dir / "Table_S2_Panel_A_dataset_construction_flow.csv", index=False
    )
    supplementary_tables["S2B"].to_csv(
        args.output_dir / "Table_S2_Panel_B_sample_structure.csv", index=False
    )
    supplementary_tables["S3"].to_csv(
        args.output_dir / "Table_S3_aggregate_prevalence.csv", index=False
    )
    supplementary_tables["S4"].to_csv(
        args.output_dir / "Table_S4_rule_level_pooled_vs_product.csv", index=False
    )
    supplementary_tables["S5"].to_csv(
        args.output_dir / "Table_S5_portal_specific_prevalence.csv", index=False
    )
    create_supplementary_docx(
        supplementary_tables,
        args.output_dir / "Tables_S2_to_S5.docx",
    )

    qa = run_qa(
        df,
        long_df,
        products,
        args.rule_summary_csv,
        supplementary_tables,
        product_key_audit,
    )
    qa.to_csv(args.output_dir / "qa_checks.csv", index=False)

    write_summary(
        long_df,
        sample_structure,
        within_product,
        cross_overlap,
        args.output_dir / "calculation_summary.txt",
    )

    failed = qa.loc[~qa["passed"]]
    print(f"Input variants: {len(df):,}")
    print(f"Distinct products: {len(products):,}")
    print(f"Outputs written to: {args.output_dir}")
    print(f"QA checks passed: {int(qa['passed'].sum())}/{len(qa)}")
    if not failed.empty:
        print(failed.to_string(index=False))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
